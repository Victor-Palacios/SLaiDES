"""Phase 8 acceptance — the agency-agents orchestrated pipeline worked example.

Covers the two PROGRESS.md "Phase 8" acceptance criteria, automated and with
**zero screenshot/render calls** (the Deck Builder contract forbids rendering a
deck to images for QA):

  (a) Feed the Visual Storyteller outline fixture — as transcribed by the Deck
      Builder into ``deck.yaml`` — through ``slidekit build``. Assert exit 0,
      an empty lint-error list, and that no external render tool is shelled out.

  (b) End-to-end: build the worked-example deck from the Brand Guardian theme
      fixture, reopen the .pptx, and assert the emitted RGB/fonts equal the
      theme block exactly (round-trip test style, extended to colors + fonts).

The fixtures live in ``examples/agency-pipeline-demo/``:
  brand-guardian-theme.md   — Brand Guardian → theme block
  storyteller-outline.md    — Visual Storyteller → slide outline
  deck.yaml                 — the Deck Builder's transcription of both
  deck.pptx                 — committed build output
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from slidekit.ir import load
from slidekit.layout import resolve
from slidekit.lint import lint
from slidekit.emit.pptx_emitter import emit_pptx, _PPTX_FONT_NAME

DEMO_DIR = Path(__file__).parent.parent.parent / "examples" / "agency-pipeline-demo"
DECK_YAML = DEMO_DIR / "deck.yaml"
OUTLINE_MD = DEMO_DIR / "storyteller-outline.md"
THEME_MD = DEMO_DIR / "brand-guardian-theme.md"


def _theme_block() -> dict:
    """The raw theme dict straight from deck.yaml (the Brand Guardian transcription)."""
    return yaml.safe_load(DECK_YAML.read_text())["theme"]


# --------------------------------------------------------------------------- #
# (a) outline fixture -> slidekit build: exit 0, empty lint errors, no render
# --------------------------------------------------------------------------- #
class TestOutlineBuildsCleanWithoutRendering:
    def test_cli_build_exits_zero_with_no_lint_errors(self, tmp_path):
        """`slidekit build` on the transcribed outline exits 0 and prints no
        lint-error JSON (build emits one JSON object per E_ error on failure)."""
        out = tmp_path / "deck.pptx"
        proc = subprocess.run(
            ["slidekit", "build", str(DECK_YAML), "-o", str(out)],
            capture_output=True,
            text=True,
        )
        assert proc.returncode == 0, proc.stderr + proc.stdout
        assert out.exists()
        # On failure each error is a JSON line with a "code"; assert none present.
        error_lines = [
            ln for ln in proc.stdout.splitlines()
            if ln.strip().startswith("{") and '"code"' in ln
        ]
        assert error_lines == [], f"unexpected lint errors: {error_lines}"
        assert "lint-clean" in proc.stdout

    def test_lint_error_list_is_empty(self):
        """The structured lint pass over the deck yields zero E_ errors."""
        deck = load(DECK_YAML)
        rd = resolve(deck)
        issues = lint(deck, rd)
        errors = [i for i in issues if i.code.startswith("E_")]
        assert errors == [], f"expected lint-clean, got {[e.code for e in errors]}"

    def test_build_shells_out_to_no_render_tool(self, tmp_path, monkeypatch):
        """The build path (load -> resolve -> lint -> emit) must not invoke any
        external process — no soffice, no pdftoppm, no screenshot tooling."""
        def _forbidden(*args, **kwargs):  # pragma: no cover - only on regression
            raise AssertionError(f"build attempted to shell out: {args!r}")

        monkeypatch.setattr(subprocess, "run", _forbidden)
        monkeypatch.setattr(subprocess, "Popen", _forbidden)
        monkeypatch.setattr(subprocess, "call", _forbidden, raising=False)
        monkeypatch.setattr(subprocess, "check_output", _forbidden, raising=False)

        deck = load(DECK_YAML)
        rd = resolve(deck)
        assert [i for i in lint(deck, rd) if i.code.startswith("E_")] == []
        emit_pptx(deck, rd, tmp_path / "deck.pptx")  # must not raise

    def test_outline_fixture_matches_deck_components(self):
        """The deck is a faithful transcription of the outline: the components
        the Storyteller suggested, in order, are the components built."""
        outline_components = re.findall(
            r"^- component:\s*(\S+)", OUTLINE_MD.read_text(), re.MULTILINE
        )
        deck = load(DECK_YAML)
        deck_components = [s.component for s in deck.slides]
        assert outline_components == deck_components
        assert len(deck_components) == 4


# --------------------------------------------------------------------------- #
# (b) theme block round-trips into emitted colors + fonts exactly
# --------------------------------------------------------------------------- #
class TestThemeRoundTrip:
    @pytest.fixture
    def emitted(self, tmp_path):
        deck = load(DECK_YAML)
        rd = resolve(deck)
        out = emit_pptx(deck, rd, tmp_path / "deck.pptx")
        return Presentation(str(out))

    def test_surface_color_round_trips(self, emitted):
        """Every slide background equals theme.palette.surface exactly."""
        surface = _theme_block()["palette"]["surface"].lstrip("#").upper()
        for i, slide in enumerate(emitted.slides):
            rgb = str(slide.background.fill.fore_color.rgb).upper()
            assert rgb == surface, f"slide {i} bg {rgb} != surface {surface}"

    def test_font_round_trips(self, emitted):
        """Every text run uses the PowerPoint name of theme.font, nothing else."""
        theme_font = _theme_block()["font"]
        expected = _PPTX_FONT_NAME[theme_font]
        seen = {
            run.font.name
            for slide in emitted.slides
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            for para in shape.text_frame.paragraphs
            for run in para.runs
        }
        assert seen == {expected}, f"fonts {seen} != {{{expected!r}}}"

    def test_text_colors_round_trip_to_theme_roles(self, emitted):
        """Every text run color is exactly a theme role and nothing else. Body text
        -> text, captions/page-numbers -> muted, emphasised headings -> primary, hero
        figures -> accent (Phase 10 visual-polish pass). No raw colors leak in."""
        palette = _theme_block()["palette"]
        roles = {palette[r].lstrip("#").upper() for r in ("text", "muted", "primary", "accent")}
        text = palette["text"].lstrip("#").upper()
        seen = {
            str(run.font.color.rgb).upper()
            for slide in emitted.slides
            for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX
            for para in shape.text_frame.paragraphs
            for run in para.runs
        }
        assert seen <= roles, f"non-theme colors emitted: {seen - roles}"
        assert text in seen, "theme text color never emitted"
