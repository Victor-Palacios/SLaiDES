"""Slide-layout structure of the emitted pptx package.

Every other emitter test goes through python-pptx's object model, which is exactly
why the original defect was invisible: the package looked fine object-wise while
every slide pointed at the stock "Blank" layout, so Google Slides' Slide > Edit
theme listed only the 11 Office layouts and none of the slidekit designs. These
tests read the raw OOXML parts and relationships instead.
"""
from __future__ import annotations

import re
import typing
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation

from slidekit.emit.prototypes import PROTOTYPES
from slidekit.emit.pptx_emitter import _layout_name, emit_pptx
from slidekit.ir import load
from slidekit.ir.models import Slide
from slidekit.layout import resolve

EXAMPLES_DIR = Path(__file__).parent.parent.parent / "examples"

_LAYOUT_PART_RE = re.compile(r"^ppt/slideLayouts/slideLayout\d+\.xml$")
_STOCK_LAYOUT_COUNT = 11


@pytest.fixture(scope="module")
def emitted(tmp_path_factory) -> zipfile.ZipFile:
    """A five-slide, five-component deck emitted once for the whole module."""
    deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
    out = emit_pptx(deck, resolve(deck), tmp_path_factory.mktemp("pptx") / "d.pptx")
    return zipfile.ZipFile(out)


def _layout_parts(z: zipfile.ZipFile) -> list[str]:
    return [n for n in z.namelist() if _LAYOUT_PART_RE.match(n)]


def _cSld_name(z: zipfile.ZipFile, part: str) -> str:
    xml = z.read(part).decode()
    m = re.search(r"<p:cSld[^>]*\bname=\"([^\"]*)\"", xml)
    return m.group(1) if m else ""


def _layout_for_slide(z: zipfile.ZipFile, one_based: int) -> str:
    rels = z.read(f"ppt/slides/_rels/slide{one_based}.xml.rels").decode()
    return re.search(r"slideLayout\d+\.xml", rels).group(0)


class TestPrototypeCoverage:
    def test_every_component_has_prototype_art(self):
        """A new component must not silently miss its layout."""
        models = typing.get_args(typing.get_args(Slide)[0])
        components = {
            typing.get_args(m.model_fields["component"].annotation)[0] for m in models
        }
        assert components - set(PROTOTYPES) == set(), (
            "components with no prototype — run scripts/build_prototypes.py"
        )


class TestLayoutParts:
    def test_one_layout_per_component_plus_stock(self, emitted):
        assert len(_layout_parts(emitted)) == _STOCK_LAYOUT_COUNT + len(PROTOTYPES)

    def test_each_design_named_exactly_once(self, emitted):
        names = [_cSld_name(emitted, p) for p in _layout_parts(emitted)]
        for component in PROTOTYPES:
            assert names.count(_layout_name(component)) == 1, (
                f"{component}: expected exactly one layout named "
                f"{_layout_name(component)!r}, got {names.count(_layout_name(component))}"
            )

    def test_layout_names_are_unique(self, emitted):
        """Our names must not collide with the stock Office ones."""
        names = [_cSld_name(emitted, p) for p in _layout_parts(emitted)]
        assert len(names) == len(set(names)), "duplicate layout names"

    def test_stock_office_layouts_are_kept(self, emitted):
        names = {_cSld_name(emitted, p) for p in _layout_parts(emitted)}
        assert {"Blank", "Title Slide", "Title and Content"} <= names

    def test_layouts_carry_art(self, emitted):
        """A layout with no shapes would show up empty in the theme editor."""
        by_name = {_cSld_name(emitted, p): p for p in _layout_parts(emitted)}
        for component in PROTOTYPES:
            xml = emitted.read(by_name[_layout_name(component)]).decode()
            assert "<p:sp>" in xml or "<p:pic>" in xml, (
                f"{component}: layout carries no shapes"
            )

    def test_layouts_are_preserved(self, emitted):
        """preserve=1 stops PowerPoint dropping layouts no slide happens to use."""
        for part in _layout_parts(emitted):
            if _cSld_name(emitted, part).startswith("slidekit"):
                assert 'preserve="1"' in emitted.read(part).decode()


class TestMasterRegistration:
    def test_master_lists_every_layout(self, emitted):
        master = emitted.read("ppt/slideMasters/slideMaster1.xml").decode()
        ids = re.findall(r"<p:sldLayoutId\b[^>]*/>", master)
        assert len(ids) == _STOCK_LAYOUT_COUNT + len(PROTOTYPES)

    def test_layout_ids_are_unique(self, emitted):
        master = emitted.read("ppt/slideMasters/slideMaster1.xml").decode()
        ids = re.findall(r'<p:sldLayoutId[^>]*\bid="(\d+)"', master)
        assert len(ids) == len(set(ids))

    def test_every_layout_rid_resolves(self, emitted):
        master = emitted.read("ppt/slideMasters/slideMaster1.xml").decode()
        rels = emitted.read("ppt/slideMasters/_rels/slideMaster1.xml.rels").decode()
        targets = dict(re.findall(r'Id="([^"]+)"[^>]*Target="([^"]+)"', rels))
        for rid in re.findall(r'<p:sldLayoutId[^>]*r:id="([^"]+)"', master):
            assert rid in targets, f"{rid} has no relationship"
            assert "slideLayout" in targets[rid]

    def test_content_types_declare_every_layout(self, emitted):
        ct = emitted.read("[Content_Types].xml").decode()
        for part in _layout_parts(emitted):
            assert "/" + part in ct, f"{part} missing a content-type override"


class TestSlideBinding:
    def test_slide_uses_layout_for_its_component(self, emitted):
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        for i, rs in enumerate(resolve(deck).slides, start=1):
            part = "ppt/slideLayouts/" + _layout_for_slide(emitted, i)
            assert _cSld_name(emitted, part) == _layout_name(rs.component)

    def test_slides_no_longer_share_the_blank_layout(self, emitted):
        """The actual regression: every slide used to point at slideLayout7 (Blank)."""
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        n = len(resolve(deck).slides)
        used = {_layout_for_slide(emitted, i) for i in range(1, n + 1)}
        assert len(used) == n, "distinct components must not share one layout"

    def test_slides_suppress_inherited_layout_graphics(self, emitted):
        """Layout art is for the theme editor; it must not paint onto slides."""
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        for i in range(1, len(resolve(deck).slides) + 1):
            xml = emitted.read(f"ppt/slides/slide{i}.xml").decode()
            assert re.search(r"<p:sld\b[^>]*showMasterSp=\"0\"", xml)


class TestPresentationMetadata:
    def test_slide_size_type_matches_dimensions(self, emitted):
        """The 4:3 default template left a stale type on a 16:9 canvas."""
        xml = emitted.read("ppt/presentation.xml").decode()
        sldSz = re.search(r"<p:sldSz[^>]*/>", xml).group(0)
        assert 'cx="12192000"' in sldSz and 'type="screen16x9"' in sldSz


class TestLayoutsAreEditable:
    """A layout you can pick but not type into is a picture, not a template.

    Plain shapes on a layout are decoration: they render on a slide using the
    layout but cannot be selected or edited there. Only placeholders get
    instantiated as editable content, so every design's text must be one.
    """

    def _slidekit_layouts(self, z):
        return {
            _cSld_name(z, p): z.read(p).decode()
            for p in _layout_parts(z)
            if _cSld_name(z, p).startswith("slidekit")
        }

    def test_every_design_exposes_editable_text(self, emitted):
        for name, xml in self._slidekit_layouts(emitted).items():
            assert re.findall(r"<p:ph\b[^>]*/>", xml), (
                f"{name}: no placeholders — a slide using it could not be edited"
            )

    def test_at_most_one_title_per_layout(self, emitted):
        """OOXML allows a single title placeholder per layout."""
        for name, xml in self._slidekit_layouts(emitted).items():
            titles = re.findall(r'<p:ph\b[^>]*type="title"', xml)
            assert len(titles) <= 1, f"{name}: {len(titles)} title placeholders"

    def test_placeholder_indices_are_unique(self, emitted):
        for name, xml in self._slidekit_layouts(emitted).items():
            idxs = re.findall(r'<p:ph\b[^>]*\bidx="(\d+)"', xml)
            assert len(idxs) == len(set(idxs)), f"{name}: duplicate placeholder idx"

    def test_placeholders_avoid_reserved_indices(self, emitted):
        """10/11/12 belong to the date, footer and slide-number placeholders."""
        for name, xml in self._slidekit_layouts(emitted).items():
            idxs = {int(i) for i in re.findall(r'<p:ph\b[^>]*\bidx="(\d+)"', xml)}
            assert not (idxs & {10, 11, 12}), f"{name}: reuses a reserved idx"

    def test_emitted_slides_carry_no_placeholders(self, emitted):
        """add_slide clones layout placeholders; our slides draw their own content,
        so those clones would be empty strays."""
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        for i in range(1, len(resolve(deck).slides) + 1):
            xml = emitted.read(f"ppt/slides/slide{i}.xml").decode()
            assert not re.findall(r"<p:ph\b[^>]*/>", xml), f"slide{i}: stray placeholder"

    def test_new_slide_from_design_is_editable(self, tmp_path):
        """End to end: pick a design's layout, get editable text slots."""
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        out = emit_pptx(deck, resolve(deck), tmp_path / "d.pptx")

        prs = Presentation(str(out))
        layout = prs.slide_layouts.get_by_name(_layout_name("pyramid"))
        slide = prs.slides.add_slide(layout)

        placeholders = list(slide.placeholders)
        assert len(placeholders) >= 2, "a picked design must expose editable slots"
        placeholders[1].text_frame.text = "edited"
        assert placeholders[1].text_frame.text == "edited"


class TestShapesAreSelectable:
    """The design's shapes must be editable too, not just its text.

    A plain shape on a layout renders on the slide but cannot be selected, so its
    color and size cannot be changed. Promoting it to a placeholder makes the
    clone on the slide a real, selectable shape.
    """

    def test_shape_bearing_designs_promote_their_shapes(self, emitted):
        """nested-circles is three ellipses plus text — all must be placeholders."""
        by_name = {_cSld_name(emitted, p): p for p in _layout_parts(emitted)}
        xml = emitted.read(by_name[_layout_name("nested-circles")]).decode()
        n_ph = len(re.findall(r"<p:ph\b[^>]*/>", xml))
        n_sp = xml.count("<p:sp>")
        assert n_ph == n_sp, (
            f"{n_sp - n_ph} shape(s) left as unselectable decoration"
        )

    def test_no_theme_style_overrides_design_colors(self, emitted):
        """<p:style> points at theme accent1 and would beat our literal fill on the
        cloned shape, turning the design stock-Office blue."""
        for part in _layout_parts(emitted):
            if _cSld_name(emitted, part).startswith("slidekit"):
                assert "<p:style>" not in emitted.read(part).decode()

    def test_shapes_keep_a_text_body(self, emitted):
        """Dropping the empty <p:txBody> renders stray glyph marks at shape edges."""
        by_name = {_cSld_name(emitted, p): p for p in _layout_parts(emitted)}
        xml = emitted.read(by_name[_layout_name("nested-circles")]).decode()
        assert xml.count("<p:txBody>") == xml.count("<p:sp>")

    def test_new_slide_exposes_selectable_shapes(self, tmp_path):
        """End to end: pick nested-circles, get shapes you can recolor and resize."""
        deck = load(EXAMPLES_DIR / "09_full_deck.yaml")
        out = emit_pptx(deck, resolve(deck), tmp_path / "d.pptx")

        prs = Presentation(str(out))
        layout = prs.slide_layouts.get_by_name(_layout_name("nested-circles"))
        slide = prs.slides.add_slide(layout)

        # Three circles + title + three value/label pairs.
        assert len(list(slide.placeholders)) == len(
            [s for s in layout.shapes if s.shape_type is not None]
        )
        shape = list(slide.placeholders)[-1]
        shape.width, shape.height = 1234567, 1234567
        assert shape.width == 1234567, "a picked design's shape must be resizable"
