"""Per-slide speaker notes.

Notes are presenter-only: the pptx emitter writes them to the slide's notes page
(what PowerPoint's and Google Slides' speaker-notes pane reads). They must never
appear on the slide itself, must not affect geometry or lint, and must not be
serialized into the layout JSON (they are prose, not geometry, so they would
otherwise churn every golden).
"""
from __future__ import annotations

import re

from pptx import Presentation

from slidekit.ir.parse import loads
from slidekit.layout import resolve
from slidekit.lint.checks import lint, has_errors
from slidekit.emit.pptx_emitter import emit_pptx

NOTE = "Open by naming the problem out loud. Do not advance until they answer."

_DECK = f"""
version: 1
slides:
  - component: title-slide
    title: Has Notes
    notes: "{NOTE}"
  - component: bullet-list
    title: No Notes
    items: [alpha, beta]
"""


def _resolve(text: str = _DECK):
    deck = loads(text)
    return deck, resolve(deck)


# ── IR ────────────────────────────────────────────────────────────────────────

def test_notes_are_optional_and_parsed():
    deck, _ = _resolve()
    assert deck.slides[0].notes == NOTE
    assert deck.slides[1].notes is None


def test_notes_available_on_every_component():
    """`notes` lives on the shared slide base, so any component accepts it."""
    deck = loads(
        "version: 1\nslides:\n"
        "  - component: code\n    title: T\n    code: |\n      x = 1\n"
        '    notes: "explain the loop"\n'
        "  - component: table-slide\n    title: T\n    headers: [A, B]\n"
        '    rows: [[A, B]]\n    notes: "read the third row aloud"\n'
    )
    assert [s.notes for s in deck.slides] == ["explain the loop", "read the third row aloud"]


# ── notes must not leak into geometry ─────────────────────────────────────────

def test_notes_reach_the_resolved_slide_but_not_the_nodes():
    _, rd = _resolve()
    assert rd.slides[0].notes == NOTE
    assert rd.slides[1].notes is None
    # The note text must not become a drawn node on the slide.
    for n in rd.slides[0].all_nodes():
        assert NOTE not in (n.text_content or "")


def test_notes_are_not_serialized_into_layout_json():
    """Prose in the golden JSON would churn every deck's snapshot."""
    _, rd = _resolve()
    assert "notes" not in rd.slides[0].to_dict()
    assert NOTE not in rd.to_json()


def test_notes_do_not_affect_lint_or_geometry():
    deck_with, rd_with = _resolve()
    assert not has_errors(lint(deck_with, rd_with))
    # Identical deck minus the notes resolves to identical geometry. node_id comes
    # from a process-global counter, so it differs between any two resolves —
    # strip it exactly as the golden tests do (tests/test_layout/test_golden.py).
    _, rd_without = _resolve(_DECK.replace(f'    notes: "{NOTE}"\n', ""))
    strip = lambda rd: re.sub(r'"node_id": "[^"]*"', '"node_id": ""', rd.to_json())
    assert strip(rd_with) == strip(rd_without)


# ── pptx notes page ───────────────────────────────────────────────────────────

def test_pptx_writes_notes_to_the_notes_page(tmp_path):
    deck, rd = _resolve()
    out = tmp_path / "d.pptx"
    emit_pptx(deck, rd, out)
    prs = Presentation(str(out))
    first, second = prs.slides[0], prs.slides[1]
    assert first.has_notes_slide
    assert first.notes_slide.notes_text_frame.text == NOTE
    # A slide with no notes gets no notes part at all.
    assert not second.has_notes_slide


def test_note_text_is_not_drawn_on_the_slide(tmp_path):
    deck, rd = _resolve()
    out = tmp_path / "d.pptx"
    emit_pptx(deck, rd, out)
    prs = Presentation(str(out))
    on_slide = " ".join(
        s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame
    )
    assert NOTE not in on_slide
