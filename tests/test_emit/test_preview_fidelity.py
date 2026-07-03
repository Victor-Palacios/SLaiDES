"""Preview/emitter fidelity guards (operator feedback FB-021 / FB-023, 2026-07-02).

The HTML preview used to silently DROP box nodes and per-node text colours: accent
bars, KPI rules, bar-chart bars, the this-vs-that badge and the code slide's syntax
colours simply never appeared on the feedback website, so the operator was reviewing
slides that looked nothing like the real .pptx/.pdf output. These tests pin the
contract: every node type and colour/alignment attribute a layout emits must survive
into the preview HTML, and paragraph alignment must reach the pptx emitter.
"""
from pathlib import Path

from slidekit.emit.html_preview import _node_html
from slidekit.ir import load
from slidekit.layout import resolve
from slidekit.layout.models import Rect, ResolvedNode

ROOT = Path(__file__).resolve().parent.parent.parent
_PALETTE = load(ROOT / "examples" / "21_code.yaml").theme.palette


def test_box_nodes_render_in_html_preview():
    node = ResolvedNode("b1", "box", Rect(0, 0, 914400, 914400), fill_color="#123456")
    html = _node_html(node, _PALETTE)
    assert "node-box" in html
    assert "background:#123456" in html


def test_text_color_override_reaches_html_preview():
    node = ResolvedNode("t1", "text", Rect(0, 0, 914400, 914400),
                        text_content="x", text_color="#E6EDF3", size_pt=34.0)
    assert "color:#E6EDF3" in _node_html(node, _PALETTE)


def test_center_alignment_reaches_html_preview():
    node = ResolvedNode("t2", "text", Rect(0, 0, 914400, 914400),
                        text_content="x", size_pt=34.0, align="center")
    assert "text-align:center" in _node_html(node, _PALETTE)
    # and left stays the default (no stray center)
    node.align = None
    assert "text-align:center" not in _node_html(node, _PALETTE)


def test_code_slide_preview_carries_syntax_colours_and_dots():
    deck = load(ROOT / "examples" / "21_code.yaml")
    rd = resolve(deck)
    html = "\n".join(_node_html(n, deck.theme.palette)
                     for n in rd.slides[0].nodes + rd.slides[0].chrome)
    from slidekit.layout.engine import _CODE_COMMENT, _CODE_FUNC, _CODE_STRING
    assert f"color:{_CODE_COMMENT}" in html  # comment lines
    assert f"color:{_CODE_STRING}" in html   # string literals
    assert f"color:{_CODE_FUNC}" in html     # function names
    assert html.count("node-box") >= 3       # the three window dots


def test_pptx_alignment_follows_node_align():
    from pptx.enum.text import PP_ALIGN
    from pptx import Presentation
    from slidekit.emit.pptx_emitter import _emit_text

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    centered = ResolvedNode("t3", "text", Rect(0, 0, 914400, 914400),
                            text_content="x", size_pt=34.0, align="center")
    left = ResolvedNode("t4", "text", Rect(0, 0, 914400, 914400),
                        text_content="x", size_pt=34.0)
    _emit_text(slide, centered, _PALETTE)
    _emit_text(slide, left, _PALETTE)
    aligns = [sh.text_frame.paragraphs[0].alignment for sh in slide.shapes]
    assert aligns == [PP_ALIGN.CENTER, PP_ALIGN.LEFT]
