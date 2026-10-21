"""E_COLOR_COUNT — the operator's three-colour rule (2026-09-05, enforced 2026-10-21).

A slide may paint at most three colours, counting its background; chrome (page
number, Source link) and image pixels do not count. The check shares the emitters'
own colour logic (slidekit.emit.colors), and the last test proves that by reading
back the colours a real .pptx was written with.
"""
from __future__ import annotations

import textwrap

from pptx import Presentation
from pptx.dml.color import RGBColor

from slidekit.emit.colors import slide_colors, text_hex
from slidekit.emit.pptx_emitter import emit_pptx
from slidekit.ir.models import Palette
from slidekit.ir.parse import loads
from slidekit.layout import resolve
from slidekit.lint.checks import MAX_SLIDE_COLOURS, lint


def _codes(text: str) -> list[str]:
    deck = loads(textwrap.dedent(text))
    return [e.code for e in lint(deck, resolve(deck)) if e.severity == "error"]


VS = """\
version: 1
{palette}
slides:
  - component: this-vs-that
    title: "Why People Use It"
    left: {{value: "Server", label: "You build and run it"}}
    right: {{value: "Supabase", label: "You set up a project"}}
"""


def test_default_palette_is_exactly_three_colours():
    p = Palette()
    assert {p.primary, p.accent} == {"#0072CE"}
    assert p.muted == p.text
    assert len({p.surface, p.accent, p.text}) == MAX_SLIDE_COLOURS


def test_versus_caption_renders_in_ink_not_grey():
    deck = loads(VS.format(palette=""))
    slide = resolve(deck).slides[0]
    vs = [n for n in slide.nodes if (n.text_content or "").strip().lower() == "vs"]
    assert vs, "this-vs-that should draw a 'vs' caption"
    assert text_hex(vs[0], deck.theme.palette).upper() == deck.theme.palette.text.upper()
    assert "E_COLOR_COUNT" not in _codes(VS.format(palette=""))


def test_a_grey_muted_role_is_a_fourth_colour():
    grey = 'theme:\n  palette: {muted: "#8A8A9A"}'
    assert "E_COLOR_COUNT" in _codes(VS.format(palette=grey))


def test_a_second_accent_is_a_fourth_colour():
    two_blues = textwrap.dedent("""\
        version: 1
        theme:
          palette: {primary: "#123A6B"}
        slides:
          - component: two-panel-list
            title: "Two Panels"
            left: {title: "Before", items: ["One thing"]}
            right: {title: "After", items: ["Another"]}
        """)  # no accent_side: the left panel takes `primary`, a second blue
    assert "E_COLOR_COUNT" in _codes(two_blues)


def test_screenshot_pixels_do_not_count():
    # A full-colour screenshot is content shown as captured, not a palette choice.
    assert "E_COLOR_COUNT" not in _codes("""\
        version: 1
        slides:
          - component: image-half-bleed
            title: "A Screenshot"
            image: {type: image, path: "examples/assets/supabase/home.png"}
            content:
              - {type: text, content: "Shown as captured"}
        """)


def test_chrome_does_not_count():
    # A grey page number is chrome: it may differ from the slide's three colours.
    # The bullet list paints white, accent title and ink items — three — and the
    # grey lives only on the page number, so the slide must still pass.
    text = textwrap.dedent("""\
        version: 1
        theme:
          palette: {muted: "#8A8A9A"}
        page_numbers: {enabled: true}
        slides:
          - component: bullet-list
            title: "Plain List"
            items: ["One", "Two"]
        """)
    deck = loads(text)
    slide = resolve(deck).slides[0]
    chrome = [n for n in slide.all_nodes() if n.is_chrome and n.node_type == "text"]
    assert chrome and text_hex(chrome[0], deck.theme.palette) == "#8A8A9A"
    assert "#8A8A9A" not in slide_colors(slide, deck.theme.palette)
    assert "E_COLOR_COUNT" not in _codes(text)


def _pptx_colours(slide) -> set[str]:
    """Every colour a pptx slide was written with: background, shape fills, text runs."""
    seen: set[str] = set()
    bg = slide.background.fill
    if bg.type is not None:
        seen.add(str(bg.fore_color.rgb))
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip() and run.font.color and run.font.color.type is not None:
                        seen.add(str(run.font.color.rgb))
        try:
            if shape.fill.type is not None and isinstance(shape.fill.fore_color.rgb, RGBColor):
                seen.add(str(shape.fill.fore_color.rgb))
        except (AttributeError, TypeError, ValueError):
            pass
    return {f"#{c}".upper() for c in seen}


def test_lint_counts_exactly_what_the_pptx_emitter_paints(tmp_path):
    deck = loads(textwrap.dedent("""\
        version: 1
        slides:
          - component: this-vs-that
            title: "Why People Use It"
            left: {value: "Server", label: "You build and run it"}
            right: {value: "Supabase", label: "You set up a project"}
          - component: two-panel-list
            title: "Two Panels"
            left: {title: "Before", items: ["One thing", "Two things"]}
            right: {title: "After", items: ["Another", "And more"]}
          - component: code
            title: "example.py"
            code: |
              def greet(name):
                  return "hi " + name
        """))
    rd = resolve(deck)
    out = tmp_path / "d.pptx"
    emit_pptx(deck, rd, out)
    prs = Presentation(str(out))
    for rs, ps in zip(rd.slides, prs.slides):
        assert _pptx_colours(ps) == slide_colors(rs, deck.theme.palette), rs.component
        assert len(_pptx_colours(ps)) <= MAX_SLIDE_COLOURS, rs.component
